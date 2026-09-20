"""第二问的统一结果导出：一个 Excel、同名 CSV、运行摘要和四张结果图。

这是主程序的一部分，直接使用 Python 环境即可；不改变工作目录，也不打开绘图窗口。
Excel 保存已计算的分析快照。调整模型参数后，应重新运行 main.py 更新所有结果。
"""

from __future__ import annotations

import json
import math
import re
import unicodedata
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib import font_manager
import numpy as np
import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


SHEET_NAMES = {
    "pair_validation": "案件对验证",
    "threshold_validation": "关联阈值选择",
    "clustering_validation": "已知组织聚类验证",
    "score_reliability": "分数分箱检验",
    "split_audit": "训练验证划分",
    "split_events": "训练验证事件清单",
    "same_country_validation": "同国案件压力测试",
    "hard_negative_validation": "同国案件压力测试",
    "group_validation": "组级候选验证",
    "group_threshold_validation": "组级拒识阈值选择",
}
SPLIT_NAMES = {
    "train_2015": "2015年训练集",
    "validation_unseen_organizations": "独立组织验证集",
    "test_unseen_organizations": "独立组织测试集",
    "test_temporal_2016": "2016年时间测试集",
    "test_unseen_organizations_fixed_subsample": "独立组织固定子样本",
}
OUTPUT_ORDER = [
    "表2嫌疑排序", "嫌疑支持度", "前五线索组", "全部线索组", "案件分组明细",
    "相似案件证据", "pair_validation", "hard_negative_validation", "group_validation",
    "clustering_validation", "阈值敏感性",
    "事件子抽样稳定性", "拒识阈值敏感性", "危害排序口径", "threshold_validation",
    "group_threshold_validation", "score_reliability", "筛选口径核对", "旧附件差异", "eventid差异说明",
    "split_audit", "split_events",
]
COLORS = ["#245A81", "#3782A8", "#63A3B4", "#BB964C", "#B76850"]


def _font_settings() -> dict:
    """优先使用系统中文字体；找不到中文字体时由 matplotlib 的字体回退处理。"""
    available = {font.name for font in font_manager.fontManager.ttflist}
    candidates = ["PingFang SC", "Heiti TC", "Hiragino Sans GB", "Microsoft YaHei",
                  "SimHei", "Noto Sans CJK SC", "WenQuanYi Zen Hei", "Arial Unicode MS"]
    fonts = [name for name in candidates if name in available]
    return {"font.family": "sans-serif", "font.sans-serif": fonts + ["DejaVu Sans"],
            "axes.unicode_minus": False, "font.size": 10, "axes.titlesize": 13,
            "axes.labelsize": 10, "axes.spines.top": False, "axes.spines.right": False,
            "figure.facecolor": "white", "axes.facecolor": "white", "savefig.facecolor": "white"}


def _json_safe(value):
    """将 NumPy/Pandas 标量和缺失值转为标准 JSON，避免 NaN 等非标准字面量。"""
    if isinstance(value, dict):
        return {str(key): _json_safe(item) for key, item in value.items()}
    if isinstance(value, (list, tuple, set, np.ndarray, pd.Series, pd.Index)):
        return [_json_safe(item) for item in value]
    if value is None or value is pd.NA or value is pd.NaT:
        return None
    if isinstance(value, (bool, np.bool_)):
        return bool(value)
    if isinstance(value, (int, np.integer)):
        return int(value)
    if isinstance(value, (float, np.floating)):
        return float(value) if np.isfinite(value) else None
    if isinstance(value, (Path, pd.Timestamp)):
        return str(value)
    return value if isinstance(value, str) else str(value)


def _identifier(value):
    """事件编号按文本存储，保留全部位数；缺失编号继续留空。"""
    if value is None or value is pd.NA or pd.isna(value):
        return None
    if isinstance(value, (int, np.integer)):
        return str(int(value))
    if isinstance(value, (float, np.floating)) and np.isfinite(value) and value.is_integer():
        return str(int(value))
    return re.sub(r"^(\d+)\.0$", r"\1", str(value).strip())


def _is_identifier(column) -> bool:
    text = str(column).lower()
    return "eventid" in text or "事件编号" in text


def _clean_table(frame: pd.DataFrame) -> pd.DataFrame:
    """只转换显示所需的数据类型，不改排名、支持度或任何模型计算值。"""
    table = frame.copy()
    table.columns = [str(column) for column in table.columns]
    for column in table:
        if _is_identifier(column):
            table[column] = table[column].map(_identifier)
    return table.replace([np.inf, -np.inf], np.nan)


def _display_width(value) -> int:
    text = str(value)
    return max((sum(2 if unicodedata.east_asian_width(ch) in "WF" else 1 for ch in line)
                for line in text.splitlines()), default=0)


def _sheet_name(key: str, used: set[str]) -> str:
    """Excel 页签有 31 字符限制；截断后仍避免同名覆盖。"""
    base = re.sub(r"[\\/*?:\[\]]", "_", SHEET_NAMES.get(key, key))[:31] or "结果"
    name, suffix = base, 2
    while name in used:
        ending = f"_{suffix}"
        name, suffix = base[:31 - len(ending)] + ending, suffix + 1
    used.add(name)
    return name


def _style_sheet(sheet, table: pd.DataFrame, key: str) -> None:
    """表格保持单行表头便于再读取，冻结首行和编号列，长文本按内容换行。"""
    sheet.sheet_view.showGridLines = False
    sheet.sheet_view.zoomScale = 95
    sheet.freeze_panes = "B2" if len(table.columns) > 1 else "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.print_title_rows = "1:1"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    if key in {"表2嫌疑排序", "嫌疑支持度", "前五线索组"}:
        sheet.sheet_properties.tabColor = "245A81"
    elif key == "说明与来源":
        sheet.sheet_properties.tabColor = "9B9B9B"
    header_font = Font(name="Microsoft YaHei", size=10, bold=True, color="FFFFFF")
    body_font = Font(name="Microsoft YaHei", size=10, color="263442")
    header_fill = PatternFill("solid", fgColor="245A81")
    alternate_fill = PatternFill("solid", fgColor="F2F5F7")
    rule = Side(style="thin", color="E0E5EA")
    for cell in sheet[1]:
        cell.font, cell.fill = header_font, header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=rule)
    sheet.row_dimensions[1].height = 34
    widths = {}
    for index, column in enumerate(table.columns, 1):
        samples = table[column].dropna().head(250)
        longest = max([_display_width(column), *[_display_width(value) for value in samples]])
        width = max(13, min(54, longest + 3))
        if _is_identifier(column):
            width = max(width, 19)
        if key == "说明与来源":
            width = 40 if index == 1 else 100
        widths[index] = width
        sheet.column_dimensions[get_column_letter(index)].width = width
    # 字符串强制为文字，既避免编号科学记数，也避免来源文字被解释成公式。
    for row in sheet.iter_rows(min_row=2):
        max_lines = 1
        for cell in row:
            cell.font = body_font
            if cell.row % 2 == 0:
                cell.fill = alternate_fill
            column = table.columns[cell.column - 1]
            if isinstance(cell.value, str):
                cell.data_type = "s"
                width = widths[cell.column]
                wrap = _display_width(cell.value) > width - 2 or "\n" in cell.value
                cell.alignment = Alignment(horizontal="left", vertical="top" if wrap else "center",
                                           wrap_text=wrap, indent=0 if _is_identifier(column) else 1)
                if wrap:
                    max_lines = max(max_lines, math.ceil(_display_width(cell.value) / max(8, width - 2)))
                if _is_identifier(column):
                    cell.number_format = "@"
            else:
                cell.alignment = Alignment(horizontal="right", vertical="center")
                if isinstance(cell.value, (float, int)) and not isinstance(cell.value, bool):
                    if column in {"iyear", "imonth", "iday", "attacktype1", "targtype1", "weaptype1", "claimed"}:
                        cell.number_format = "0"
                    elif "嫌疑人" in column and "支持" not in column:
                        cell.number_format = "0"
                    elif column in {"受害者死亡", "受害者受伤", "已记录受害者死亡", "已记录受害者受伤"}:
                        cell.number_format = "#,##0"
                    elif isinstance(cell.value, int) or any(word in column for word in ["事件数", "次数", "格数", "国家数", "候选数"]):
                        cell.number_format = "#,##0"
                    else:
                        cell.number_format = "0.000"
        sheet.row_dimensions[row[0].row].height = 20 if max_lines == 1 else min(300, 15 * max_lines + 5)


def _notes_table(summary: dict) -> pd.DataFrame:
    """把解释和来源集中在最后一页；完整嵌套配置同时保存为 JSON。"""
    notes = [
        ("结果更新", "本工作簿是 Python 模型计算的结果快照；修改参数或附件后，重新运行 main.py 更新 Excel、CSV 和图表。"),
        ("表2读法", "每列固定对应一个嫌疑人代号，单元格为该事件中该候选的嫌疑排名；1表示最靠前，空白表示未达到判定条件。"),
        ("线索组与支持度", "线索组是具有相似案件特征的模型分组，可能拆分真实组织或合并不同组织；支持度不是作案概率，也不确认真实组织身份。"),
        ("危害累计", "累计危害是观测案件危害得分之和，受案件数与聚类大小影响；应同时参考平均危害、最高单次危害和原始伤亡。"),
        ("数据范围", "以主程序筛选口径核对表为准；无组织认领与组织名称未知不是同一概念。原附件2的预筛选口径另行核对。"),
        ("验证范围", "案件对指标依赖正负样本抽样比例；已知组织上的测试成绩不能作为未知案件真实归属的证明。"),
        ("原始资料", "2018年中国研究生数学建模竞赛C题，任务2；本项目附件1.xlsx、附件2.xlsx、附件3.xlsx。"),
        ("GTD字段说明", "https://www.start.umd.edu/sites/default/files/2024-10/Codebook.pdf"),
        ("分组验证依据", "https://scikit-learn.org/stable/modules/generated/sklearn.model_selection.GroupShuffleSplit.html"),
        ("聚类评估依据", "https://scikit-learn.org/stable/modules/clustering.html#clustering-performance-evaluation"),
    ]
    def collect(data, prefix=""):
        for key, value in data.items():
            label = f"{prefix} / {key}" if prefix else str(key)
            if isinstance(value, dict):
                collect(value, label)
            else:
                safe = _json_safe(value)
                text = json.dumps(safe, ensure_ascii=False) if isinstance(safe, list) else str(safe)
                # 运行日志中的绝对路径不必出现在面向阅读的说明表中。
                if text.startswith("/"):
                    text = Path(text).name
                notes.append((label, text))
    collect(summary)
    return pd.DataFrame(notes, columns=["项目", "说明或数值"])


def _no_data(ax, title: str, message="本次没有可绘制的数据"):
    ax.set_title(title, loc="left")
    ax.text(.5, .5, message, transform=ax.transAxes, ha="center", va="center", color="#666666")
    ax.set_axis_off()


def _save_figure(fig, path: Path):
    fig.savefig(path, dpi=180, bbox_inches="tight")
    plt.close(fig)


def _plot_top_groups(tables, directory):
    top = tables.get("前五线索组", pd.DataFrame()).head(5)
    fig, axes = plt.subplots(1, 2, figsize=(11, 4.7), constrained_layout=True)
    for ax, column, title, unit in zip(axes, ["累计危害", "事件数"], ["累计危害", "案件数量"], ["累计危害指数", "起"]):
        if top.empty or column not in top:
            _no_data(ax, title)
            continue
        labels = top["嫌疑人代号"].astype(str) if "嫌疑人代号" in top else pd.Series([f"{i + 1}号" for i in range(len(top))])
        labels = [label if "号" in label else f"{label}号" for label in labels]
        values = pd.to_numeric(top[column], errors="coerce").to_numpy()
        bars = ax.barh(labels, values, color=COLORS[:len(top)], height=.58)
        ax.invert_yaxis()
        ax.set_title(title, loc="left")
        ax.set_xlabel(unit)
        ax.grid(axis="x", alpha=.18)
        ax.set_axisbelow(True)
        ax.margins(x=.18)
        ax.bar_label(bars, labels=[f"{v:,.0f}" if column == "事件数" else f"{v:,.2f}" for v in values], padding=5, fontsize=9)
    fig.suptitle("前五个潜在线索组", fontsize=15)
    _save_figure(fig, directory / "top5_groups.png")


def _plot_support(tables, directory):
    frame = tables.get("嫌疑支持度", pd.DataFrame())
    columns = [f"{i}号支持度" for i in range(1, 6)]
    fig, ax = plt.subplots(figsize=(10.5, max(4, .55 * len(frame) + 2)), constrained_layout=True)
    if frame.empty or not all(column in frame for column in columns):
        _no_data(ax, "表2事件的候选支持度")
    else:
        values = frame[columns].apply(pd.to_numeric, errors="coerce").to_numpy(dtype=float)
        cmap = plt.get_cmap("Blues").copy()
        cmap.set_bad("#ECEFF2")
        graphic = ax.imshow(np.ma.masked_invalid(values), cmap=cmap, vmin=0, vmax=1, aspect="auto")
        ids = frame["eventid"].map(_identifier) if "eventid" in frame else frame.index.astype(str)
        countries = frame["事件国家"].fillna("").astype(str) if "事件国家" in frame else [""] * len(frame)
        ax.set_yticks(range(len(frame)), [f"{event}  {country}".strip() for event, country in zip(ids, countries)], fontsize=9)
        ax.set_xticks(range(5), [f"{i}号" for i in range(1, 6)])
        thresholds = pd.to_numeric(frame.get("判定阈值", pd.Series(np.nan, index=frame.index)), errors="coerce").to_numpy()
        ranking = tables.get("表2嫌疑排序", pd.DataFrame())
        rank_lookup = (ranking.assign(eventid=ranking.eventid.map(_identifier)).set_index("eventid")
                       if "eventid" in ranking and ranking.eventid.is_unique else pd.DataFrame())
        for i in range(len(frame)):
            for j in range(5):
                value = values[i, j]
                rank_column = f"{j + 1}号嫌疑人"
                eventid = str(ids.iloc[i]) if isinstance(ids, pd.Series) else str(ids[i])
                rank_missing = (eventid in rank_lookup.index and rank_column in rank_lookup
                                and pd.isna(rank_lookup.at[eventid, rank_column]))
                if np.isnan(value):
                    label = "无候选"
                elif np.isfinite(thresholds[i]) and value < thresholds[i]:
                    label = f"{value:.3f}\n低于阈值"
                elif rank_missing:
                    label = f"{value:.3f}\n未列入排名"
                else:
                    label = f"{value:.3f}"
                ax.text(j, i, label, ha="center", va="center", fontsize=8,
                        color="white" if np.isfinite(value) and value >= .64 else "#263442")
        ax.set_title("表2事件的候选支持度", loc="left", pad=13)
        ax.set_xlabel("表2空白表示未达到阈值或证据条数要求；数值为模型相似支持度。", labelpad=12)
        fig.colorbar(graphic, ax=ax, fraction=.035, pad=.025, label="支持度（非作案概率）")
    _save_figure(fig, directory / "suspect_support.png")


def _plot_validation(tables, directory):
    frame = tables.get("pair_validation", pd.DataFrame())
    metrics = [("average_precision", "AP"), ("precision", "精确率"), ("recall", "召回率")]
    fig, ax = plt.subplots(figsize=(10.8, 5.3), constrained_layout=True)
    if frame.empty or "split" not in frame:
        _no_data(ax, "案件对模型的独立验证")
    else:
        existing = [(key, label) for key, label in metrics if key in frame]
        locations = np.arange(len(frame))
        width = .72 / max(len(existing), 1)
        for index, (key, label) in enumerate(existing):
            values = pd.to_numeric(frame[key], errors="coerce")
            bars = ax.bar(locations + (index - (len(existing) - 1) / 2) * width, values,
                          width, label=label, color=COLORS[index])
            ax.bar_label(bars, labels=[f"{value:.2f}" if pd.notna(value) else "" for value in values], padding=3, fontsize=8)
        labels = []
        for _, row in frame.iterrows():
            label = SPLIT_NAMES.get(str(row["split"]), str(row["split"]))
            if "model" in row and pd.notna(row["model"]):
                model_label = {"full_pair_model": "完整模型", "geography_only": "仅地理基线"}.get(
                    str(row["model"]), str(row["model"]))
                label += "\n" + model_label
            if "subset" in row and pd.notna(row["subset"]):
                label += "\n" + str(row["subset"])
            labels.append(label)
        ax.set_xticks(locations, labels, fontsize=9, rotation=15 if len(frame) > 4 else 0)
        ax.set_ylim(0, 1.16)
        ax.set_ylabel("指标值")
        ax.set_title("案件对模型的独立验证", loc="left")
        ax.legend(loc="upper right", ncol=3, frameon=False)
        ax.grid(axis="y", alpha=.18)
        ax.set_axisbelow(True)
        ax.set_xlabel("指标依赖本次案件对抽样；验证集选阈值，独立测试集报告性能。", labelpad=12)
    _save_figure(fig, directory / "pair_validation.png")


def _plot_stability(tables, directory):
    fig, axes = plt.subplots(1, 2, figsize=(11, 4.8), constrained_layout=True)
    for ax, key, x_column, title in [
        (axes[0], "阈值敏感性", "阈值", "关联阈值敏感性"),
        (axes[1], "事件子抽样稳定性", "重复次数", "事件子抽样稳定性"),
    ]:
        frame = tables.get(key, pd.DataFrame())
        if frame.empty or x_column not in frame:
            _no_data(ax, title)
            continue
        x = pd.to_numeric(frame[x_column], errors="coerce")
        used = False
        for column, color in [("ARI", COLORS[0]), ("NMI", COLORS[2])]:
            if column in frame:
                ax.plot(x, pd.to_numeric(frame[column], errors="coerce"), marker="o", ms=4,
                        label=column, color=color, lw=1.6)
                used = True
        if not used:
            _no_data(ax, title, "本次结果未包含 ARI / NMI 指标")
            continue
        ax.set_title(title, loc="left")
        ax.set_xlabel(x_column)
        ax.set_ylabel("与基准分组的一致性")
        ax.set_ylim(-.05, 1.05)
        ax.grid(alpha=.18)
        ax.legend(frameon=False)
    _save_figure(fig, directory / "stability.png")


def export_results(tables: dict[str, pd.DataFrame], summary: dict, output: Path) -> None:
    """一次导出所有结果。主程序只需调用此函数，不需要用户手工整理文件。

    保留每张输入表的全部行和列，数值不做显示性取整；Excel 格式只影响呈现。
    事件编号使用文本，NaN/NA 排名留空。图表缺少数据时明确标记而不伪造数值。
    """
    output = Path(output)
    output.mkdir(parents=True, exist_ok=True)
    directory = output / "figures"
    directory.mkdir(parents=True, exist_ok=True)
    ordered_keys = [key for key in OUTPUT_ORDER if key in tables]
    ordered_keys += [key for key in tables if key not in ordered_keys]
    cleaned = {key: _clean_table(tables[key]) for key in ordered_keys}
    safe_summary = _json_safe(summary)
    used = set()
    workbook_path = output / "第二问分析结果.xlsx"
    print(f"      正在保存 {len(cleaned)} 张分析表与 Excel 汇总……", flush=True)
    with pd.ExcelWriter(workbook_path, engine="openpyxl") as writer:
        for key, table in cleaned.items():
            name = _sheet_name(key, used)
            table.to_excel(writer, sheet_name=name, index=False, na_rep="")
            _style_sheet(writer.sheets[name], table, key)
            # UTF-8 BOM 便于 Excel 正确识别中文；文件名和原始表名一致。
            filename = re.sub(r"[\\/:*?\"<>|]", "_", key) + ".csv"
            table.to_csv(output / filename, index=False, encoding="utf-8-sig", na_rep="")
        notes = _notes_table(safe_summary)
        name = _sheet_name("说明与来源", used)
        notes.to_excel(writer, sheet_name=name, index=False)
        _style_sheet(writer.sheets[name], notes, "说明与来源")
    (output / "run_summary.json").write_text(
        json.dumps(safe_summary, ensure_ascii=False, indent=2, allow_nan=False), encoding="utf-8")
    print("      正在生成前五组危害、嫌疑支持度、模型验证与稳定性图……", flush=True)
    # rc_context 避免污染用户 Spyder 会话的全局绘图字体/颜色设置。
    with plt.rc_context(_font_settings()):
        _plot_top_groups(cleaned, directory)
        _plot_support(cleaned, directory)
        _plot_validation(cleaned, directory)
        _plot_stability(cleaned, directory)
    print(f"      Excel 与四张结果图已保存：{output.resolve()}", flush=True)
