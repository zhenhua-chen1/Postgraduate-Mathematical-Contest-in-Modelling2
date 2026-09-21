"""原始附件核验、月度面板和近三年描述统计。所有清洗都在内存中进行。"""
from __future__ import annotations

import hashlib
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

REGIONS = {1: '北美', 2: '中美洲与加勒比', 3: '南美', 4: '东亚', 5: '东南亚',
           6: '南亚', 7: '中亚', 8: '西欧', 9: '东欧', 10: '中东与北非',
           11: '撒哈拉以南非洲', 12: '大洋洲'}
USE_COLUMNS = ['eventid', 'iyear', 'imonth', 'iday', 'country', 'country_txt',
               'region', 'region_txt', 'latitude', 'longitude', 'specificity',
               'attacktype1_txt', 'targtype1_txt', 'weaptype1_txt', 'gname',
               'guncertain1', 'claimed', 'suicide', 'success', 'nkill', 'nwound',
               'nkillter', 'nwoundte', 'doubtterr', 'INT_ANY', 'motive']


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open('rb') as f:
        for block in iter(lambda: f.read(1024 * 1024), b''):
            h.update(block)
    return h.hexdigest()


def read_events(path: Path, progress=print) -> pd.DataFrame:
    """流式读取并定期报进度，避免大附件长时间没有反馈。"""
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.active
        rows = iter(ws.values)
        header = next(rows)
        positions = [header.index(c) for c in USE_COLUMNS]
        records = []
        for n, row in enumerate(rows, 1):
            if row[positions[0]] is None:
                raise ValueError(f'附件第 {n + 1} 行缺少事件编号，停止而非静默删除。')
            records.append([row[i] for i in positions])
            if n % 20000 == 0:
                progress(f'已读取 {n:,}/{ws.max_row - 1:,} 起事件')
        return pd.DataFrame(records, columns=USE_COLUMNS)
    finally:
        wb.close()


def validate_events(data: pd.DataFrame) -> pd.DataFrame:
    d = data.copy()
    if d.eventid.isna().any() or d.eventid.duplicated().any():
        raise ValueError('事件编号存在缺失或重复。')
    for col in ['eventid', 'iyear', 'imonth', 'iday', 'region']:
        v = pd.to_numeric(d[col], errors='raise')
        if v.isna().any() or (v % 1 != 0).any():
            raise ValueError(f'{col} 必须是已知整数。')
        d[col] = v.astype('int64')
    if not d.iyear.between(1998, 2017).all():
        raise ValueError('当前题目仅接受 1998—2017 年附件，防止未来数据进入预测。')
    if not d.imonth.between(1, 12).all():
        raise ValueError('存在未知或非法月份，不能把这些记录当作零事件月份。')
    if not d.region.isin(REGIONS).all():
        raise ValueError('存在未知地区，不能静默丢弃。')
    # eventid 是主键，不保证前八位等于实际日期。按专门的年月字段统计。
    d['month'] = pd.to_datetime(dict(year=d.iyear, month=d.imonth, day=1))
    if not d.iday.between(0, 31).all():
        raise ValueError('存在非法日期。iday=0 表示日未知，可以参加月度统计。')
    known = d.iday > 0
    pd.to_datetime(dict(year=d.loc[known, 'iyear'], month=d.loc[known, 'imonth'],
                        day=d.loc[known, 'iday']), errors='raise')
    d['地区'] = d.region.map(REGIONS)
    for col in ['latitude', 'longitude', 'nkill', 'nwound', 'nkillter', 'nwoundte']:
        d[col] = pd.to_numeric(d[col], errors='coerce')
    # 伤亡中的负数编码是未知；总伤亡保留袭击者，不冒充“受害者伤亡”。
    for col in ['nkill', 'nwound', 'nkillter', 'nwoundte']:
        d[col] = d[col].where(d[col] >= 0)
    d['坐标状态'] = coordinate_status(d)
    return d


def coordinate_status(d: pd.DataFrame) -> pd.Series:
    status = pd.Series('有效', index=d.index)
    status.loc[(d.latitude == 0) & (d.longitude == 0)] = '双零待核实'
    status.loc[(d.latitude.abs() > 90) | (d.longitude.abs() > 180)] = '超出范围'
    status.loc[d[['latitude', 'longitude']].isna().any(axis=1)] = '缺失'
    # 单独的纬度 0 或经度 0 合法，不能一概剔除。
    return status


def connect_scores(d: pd.DataFrame, path: Path) -> pd.DataFrame:
    scores = pd.read_csv(path, usecols=['eventid', '危害等级', '危害得分', '数据完整率'])
    if scores.eventid.duplicated().any() or set(scores.eventid) != set(d.eventid):
        raise ValueError('第一问评分与附件事件集合不一致，请先检查版本。')
    if not scores['危害得分'].between(0, 1).all() or not scores['危害等级'].isin(range(1, 6)).all():
        raise ValueError('第一问评分范围或等级非法。')
    bounds = scores.groupby('危害等级')['危害得分'].agg(['min', 'max'])
    if any(bounds.loc[k, 'min'] < bounds.loc[k + 1, 'max'] - 1e-12 for k in range(1, 5)):
        raise ValueError('等级方向错误：应当一级最严重、五级最轻。')
    return d.merge(scores, on='eventid', how='left', validate='one_to_one')


def monthly_counts(d: pd.DataFrame, start='2007-01-01', end='2017-12-01') -> pd.DataFrame:
    dates = pd.date_range(start, end, freq='MS')
    # 全球各月必须有观测；仅在这个覆盖前提下将地区未记录事件的月份补为 0。
    observed = d.groupby('month').size()
    missing = dates.difference(observed.index)
    if len(missing):
        raise ValueError(f'整月数据缺失，不能直接填零：{list(missing)}')
    p = d.groupby(['month', 'region']).size().unstack('region')
    p = p.reindex(index=dates, columns=list(REGIONS)).fillna(0).astype(int)
    p.index.name = '月份'
    p.columns = [REGIONS[k] for k in p.columns]
    if not np.array_equal(p.sum(axis=1).values, observed.reindex(dates).values):
        raise AssertionError('地区与全球计数不一致。')
    return p


def describe(d: pd.DataFrame) -> dict[str, pd.DataFrame]:
    x = d[d.iyear.between(2015, 2017)].copy()
    x['一级事件'] = x['危害等级'].eq(1).astype(int)
    x['有效坐标'] = x['坐标状态'].eq('有效').astype(int)
    def aggregate(keys):
        out = x.groupby(keys, observed=True).agg(
            事件数=('eventid', 'size'), 累计危害=('危害得分', 'sum'),
            平均危害=('危害得分', 'mean'), 一级事件数=('一级事件', 'sum'),
            有效坐标数=('有效坐标', 'sum'), 已知死亡记录数=('nkill', 'count'),
            已记录死亡人数=('nkill', lambda s: s.sum(min_count=1)),
            已知受伤记录数=('nwound', 'count'),
            已记录受伤人数=('nwound', lambda s: s.sum(min_count=1)))
        out['一级占比'] = out['一级事件数'] / out['事件数']
        out['坐标覆盖率'] = out['有效坐标数'] / out['事件数']
        out['死亡信息覆盖率'] = out['已知死亡记录数'] / out['事件数']
        return out.reset_index()
    annual = aggregate('iyear').rename(columns={'iyear': '年份'})
    annual['事件数同比'] = annual['事件数'].pct_change()
    annual['有事件国家数'] = x.groupby('iyear').country.nunique().values
    region = aggregate(['地区', 'iyear']).rename(columns={'iyear': '年份'})
    month = aggregate('month').rename(columns={'month': '月份'})
    region_month = aggregate(['month', '地区']).rename(columns={'month': '月份'})
    grid = pd.MultiIndex.from_product([pd.date_range('2015-01-01', '2017-12-01', freq='MS'),
                                      REGIONS.values()], names=['月份', '地区'])
    region_month = region_month.set_index(['月份', '地区']).reindex(grid)
    additive = ['事件数', '累计危害', '一级事件数', '有效坐标数', '已知死亡记录数', '已知受伤记录数']
    region_month[additive] = region_month[additive].fillna(0)
    # 空月平均危害/占比未知（分母为零），不可伪装为零危害事件。
    region_month = region_month.reset_index()
    level = x.groupby(['iyear', '危害等级']).size().unstack(fill_value=0).reindex(columns=range(1, 6), fill_value=0)
    level = level.stack().rename('事件数').reset_index().rename(columns={'iyear': '年份'})
    level['年度占比'] = level['事件数'] / level.groupby('年份')['事件数'].transform('sum')
    region_level = x.groupby(['地区', 'iyear', '危害等级']).size().rename('事件数').reset_index().rename(columns={'iyear': '年份'})
    region_level['地区年度占比'] = region_level['事件数'] / region_level.groupby(['地区', '年份'])['事件数'].transform('sum')
    seasonal = month[['月份', '事件数']].copy()
    seasonal['年份'] = seasonal['月份'].dt.year
    seasonal['月序号'] = seasonal['月份'].dt.month
    seasonal['日均事件数'] = seasonal['事件数'] / seasonal['月份'].dt.days_in_month
    year_days = {y: 366 if pd.Timestamp(y, 12, 31).is_leap_year else 365 for y in [2015, 2016, 2017]}
    daily = x.groupby('iyear').size() / pd.Series(year_days)
    seasonal['相对当年日均水平'] = seasonal['日均事件数'] / seasonal['年份'].map(daily)
    country = aggregate(['country_txt', 'iyear']).rename(columns={'country_txt': '国家', 'iyear': '年份'})
    factors = []
    for field, name in [('attacktype1_txt', '袭击类型'), ('targtype1_txt', '目标类型'),
                        ('weaptype1_txt', '武器类型'), ('gname', '归属记录')]:
        t = aggregate([field, 'iyear']).rename(columns={field: '类别', 'iyear': '年份'})
        t.insert(0, '维度', name)
        t['年度事件占比'] = t['事件数'] / t.groupby('年份')['事件数'].transform('sum')
        factors.append(t)
    return {'年度概况': annual, '全球月度历史': month, '地区年度统计': region,
            '地区月度面板': region_month, '五级分布': level, '国家年度统计': country,
            '地区级别分布': region_level, '季节性描述': seasonal,
            '事件特征统计': pd.concat(factors, ignore_index=True)}


def spatial_analysis(d: pd.DataFrame, nlon=180, nlat=90):
    """等面积经度×sin(纬度)网格。格点变化只描述扩张/收缩，不识别因果传播。"""
    x = d[d.iyear.between(2015, 2017) & d['坐标状态'].eq('有效')].copy()
    x['经度格'] = np.minimum(np.floor((x.longitude + 180) / 360 * nlon).astype(int), nlon - 1)
    x['纬度格'] = np.minimum(np.floor((np.sin(np.radians(x.latitude)) + 1) / 2 * nlat).astype(int), nlat - 1)
    cells = x.groupby(['iyear', '经度格', '纬度格']).agg(
        事件数=('eventid', 'size'), 累计危害=('危害得分', 'sum')).reset_index().rename(columns={'iyear': '年份'})
    cells['中心经度'] = (cells['经度格'] + .5) * 360 / nlon - 180
    cells['中心纬度'] = np.degrees(np.arcsin((cells['纬度格'] + .5) * 2 / nlat - 1))
    rows = []
    for year in [2015, 2016, 2017]:
        c = cells[cells['年份'].eq(year)]
        prev = cells[cells['年份'].eq(year - 1)]
        active = set(zip(c['经度格'], c['纬度格']))
        old = set(zip(prev['经度格'], prev['纬度格']))
        new = active - old
        counts = c['事件数'].sort_values(ascending=False)
        row = {'年份': year, '网格经度数': nlon, '网格纬度数': nlat,
               '每格面积平方公里': 4 * np.pi * 6371.0088 ** 2 / (nlon * nlat),
               '空间事件数': int(counts.sum()), '活跃格数': len(active),
               '前十格事件占比': counts.head(10).sum() / counts.sum(),
               'HHI': ((counts / counts.sum()) ** 2).sum(),
               '相邻年持续格数': len(active & old) if old else np.nan,
               '相对上年新增格数': len(new) if old else np.nan,
               '相对上年消失格数': len(old - active) if old else np.nan,
               '相邻年Jaccard': len(active & old) / len(active | old) if old else np.nan,
               '新增格事件数': int(c[[p in new for p in zip(c['经度格'], c['纬度格'])]]['事件数'].sum()) if old else np.nan}
        rows.append(row)
    return cells, pd.DataFrame(rows)
