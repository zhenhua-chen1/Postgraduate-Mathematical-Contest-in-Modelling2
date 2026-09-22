#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""2019 年 D 题第一问：三份车辆数据预处理。

在 Spyder 中打开本文件并点击绿色运行箭头即可。程序不依赖当前工作目录，
原始文件始终只读，所有新结果写入同目录的“优化结果”文件夹。
"""
from __future__ import annotations

import hashlib
import importlib.metadata
import json
import platform
import sys
from datetime import datetime
from pathlib import Path
from time import perf_counter

import numpy as np
import pandas as pd
import matplotlib as mpl
from matplotlib.backends.backend_agg import FigureCanvasAgg
from matplotlib.figure import Figure
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

BASE = Path(__file__).resolve().parent
if str(BASE) not in sys.path:
    sys.path.insert(0, str(BASE))

from preprocessing import PreprocessConfig, preprocess_vehicle_data, validate_processed_data  # noqa: E402

OUTPUT = BASE / "优化结果"
INPUTS = [BASE / f"文件{i}.xlsx" for i in range(1, 4)]


class Progress:
    """在控制台实时显示进度，并同步保存完整日志。"""

    def __init__(self, output: Path):
        self.start = perf_counter()
        self.path = output / "运行日志.txt"
        self.path.write_text("", encoding="utf-8")

    def __call__(self, message: str) -> None:
        line = f"[{datetime.now():%H:%M:%S} | {perf_counter() - self.start:6.1f}s] {message}"
        print(line, flush=True)
        with self.path.open("a", encoding="utf-8") as stream:
            stream.write(line + "\n")


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def json_text(value) -> str:
    def convert(item):
        if isinstance(item, dict):
            return {key: convert(val) for key, val in item.items()}
        if isinstance(item, (list, tuple)):
            return [convert(val) for val in item]
        if isinstance(item, (np.integer, np.floating)):
            return item.item()
        if isinstance(item, (Path, pd.Timestamp)):
            return str(item)
        if isinstance(item, float) and not np.isfinite(item):
            return None
        return item

    return json.dumps(convert(value), ensure_ascii=False, indent=2, allow_nan=False)


def read_input(path: Path, file_number: int) -> pd.DataFrame:
    """读取指定原始工作表，并保留输入列顺序。"""
    expected_sheet = f"原始数据{file_number}"
    with pd.ExcelFile(path) as book:
        if expected_sheet not in book.sheet_names:
            raise ValueError(f"{path.name} 缺少工作表“{expected_sheet}”。")
        return pd.read_excel(book, sheet_name=expected_sheet)


def style_workbook(path: Path, freeze_cell: str = "A2") -> None:
    """使用轻量样式提高可读性，不改变数据值。"""
    book = load_workbook(path)
    try:
        for sheet in book.worksheets:
            sheet.freeze_panes = freeze_cell
            sheet.auto_filter.ref = sheet.dimensions
            sheet.sheet_view.showGridLines = False
            for cell in sheet[1]:
                cell.fill = PatternFill("solid", fgColor="1F4E78")
                cell.font = Font(color="FFFFFF", bold=True, name="Arial", size=10)
                cell.alignment = Alignment(horizontal="center", vertical="center")
            for column_index, header in enumerate(sheet[1], 1):
                letter = header.column_letter
                sample = [
                    str(sheet.cell(row, column_index).value)
                    if sheet.cell(row, column_index).value is not None else ""
                    for row in range(1, min(sheet.max_row, 200) + 1)
                ]
                sheet.column_dimensions[letter].width = min(max(max(map(len, sample), default=8) + 2, 10), 28)
            if sheet.title == "处理汇总":
                sheet.column_dimensions["A"].width = 12
                for letter in "BCDEFGHI":
                    sheet.column_dimensions[letter].width = 18
                sheet.column_dimensions["J"].width = 25
                sheet.column_dimensions["K"].width = 25
            elif sheet.title == "方法与阈值":
                sheet.column_dimensions["A"].width = 14
                sheet.column_dimensions["B"].width = 72
            else:
                sheet.column_dimensions["A"].width = 22
                sheet.column_dimensions["Q"].width = 28
            sheet.row_dimensions[1].height = 24
        book.save(path)
    finally:
        book.close()


def save_processed(frame: pd.DataFrame, file_number: int, output: Path) -> Path:
    path = output / f"文件{file_number}_预处理结果.xlsx"
    sheet = f"预处理数据{file_number}"
    with pd.ExcelWriter(path, engine="openpyxl", datetime_format="yyyy-mm-dd hh:mm:ss") as writer:
        frame.to_excel(writer, sheet_name=sheet, index=False)
    style_workbook(path)
    return path


def save_summary(summary: pd.DataFrame, config: PreprocessConfig, output: Path) -> Path:
    methods = pd.DataFrame(
        [
            ("时间解析", "按原时间字段解析；重复、倒序或非法时间直接报错"),
            ("短时缺测", f"仅对恰好 {config.interpolate_gap_seconds} 秒的相邻记录补1个线性插值点"),
            ("长时缺测", f"超过 {config.interpolate_gap_seconds} 秒不插值，作为新连续段"),
            ("异常加速", f"连续1秒最大加速度 {config.max_acceleration_mps2:g} m/s²"),
            ("异常减速", f"连续1秒最大减速度 {config.max_deceleration_mps2:g} m/s²"),
            ("长低速段", f"连续速度低于 {config.low_speed_kmh:g} km/h 的记录最多保留 {config.max_low_speed_records} 条"),
            ("原始值", "原附件不覆盖；输出保留原GPS车速、处理标记和原始行号"),
        ],
        columns=["项目", "规则"],
    )
    path = output / "预处理汇总.xlsx"
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        summary.to_excel(writer, sheet_name="处理汇总", index=False)
        methods.to_excel(writer, sheet_name="方法与阈值", index=False)
    style_workbook(path)
    return path


def save_summary_figure(summary: pd.DataFrame, output: Path) -> Path:
    """绘制一张三文件预处理总览图，不改变 Spyder 当前绘图后端。"""
    figure_dir = output / "figures"
    figure_dir.mkdir(exist_ok=True)
    path = figure_dir / "01_三文件预处理结果.png"
    names = summary["文件"].tolist()
    x = np.arange(len(names))
    colors = {"原始记录": "#94A3B8", "处理后记录": "#2563EB",
              "两秒补点": "#2A9D8F", "车速修正": "#E9C46A", "长低速删除": "#E76F51"}

    with mpl.rc_context({
        "font.sans-serif": ["Arial Unicode MS", "PingFang SC", "Microsoft YaHei", "SimHei", "DejaVu Sans"],
        "axes.unicode_minus": False,
        "font.size": 10,
    }):
        figure = Figure(figsize=(13.5, 6.6), facecolor="white", layout="constrained")
        FigureCanvasAgg(figure)
        axes = figure.subplots(1, 2, gridspec_kw={"width_ratios": [1, 1.25]})
        figure.suptitle("三份车辆数据预处理结果", fontsize=16, fontweight="bold")

        ax = axes[0]
        width = 0.34
        original = summary["原始记录数"].to_numpy()
        processed = summary["处理后记录数"].to_numpy()
        bars1 = ax.bar(x - width / 2, original, width, label="原始记录", color=colors["原始记录"])
        bars2 = ax.bar(x + width / 2, processed, width, label="处理后记录", color=colors["处理后记录"])
        ax.bar_label(bars1, labels=[f"{value:,}" for value in original], padding=3, fontsize=9)
        ax.bar_label(bars2, labels=[f"{value:,}" for value in processed], padding=3, fontsize=9)
        ax.set_title("记录数变化")
        ax.set_ylabel("记录数（条）")
        ax.set_xticks(x, names)
        ax.set_ylim(0, max(original) * 1.18)
        ax.legend(frameon=False, loc="upper right")
        ax.grid(axis="y", color="#E2E8F0", linewidth=0.8)
        ax.set_axisbelow(True)

        ax = axes[1]
        labels, values, bar_colors = [], [], []
        for _, row in summary.iterrows():
            for label, column in [("两秒补点", "插值新增数"), ("车速修正", "车速修正数"),
                                  ("长低速删除", "长低速删除数")]:
                labels.append(f"{row['文件']} · {label}")
                values.append(int(row[column]))
                bar_colors.append(colors[label])
        positions = np.arange(len(labels))[::-1]
        bars = ax.barh(positions, values, color=bar_colors, height=0.68)
        ax.bar_label(bars, labels=[f"{value:,}" for value in values], padding=4, fontsize=9)
        ax.set_title("各类处理记录数")
        ax.set_xlabel("记录数（条）")
        ax.set_yticks(positions, labels)
        ax.set_xlim(0, max(values) * 1.18)
        ax.grid(axis="x", color="#E2E8F0", linewidth=0.8)
        ax.set_axisbelow(True)

        for ax in axes:
            ax.spines[["top", "right"]].set_visible(False)
            ax.ticklabel_format(axis="x" if ax is axes[1] else "y", style="plain")
        figure.text(
            0.5, 0.005,
            "规则：仅填补两秒缺口；连续1 Hz记录限制异常加减速度；连续低于10 km/h的长低速段最多保留180条",
            ha="center", color="#475569", fontsize=9,
        )
        figure.savefig(path, dpi=220, bbox_inches="tight")
    return path


def main() -> None:
    OUTPUT.mkdir(parents=True, exist_ok=True)
    log = Progress(OUTPUT)
    config = PreprocessConfig()
    summary_path = OUTPUT / "run_summary.json"
    run = {"status": "running", "started_at": datetime.now().astimezone().isoformat()}
    summary_path.write_text(json_text(run), encoding="utf-8")

    try:
        log("2019年D题第一问开始；原始附件只读，结果写入优化结果")
        log(f"当前解释器：{sys.executable}")
        missing = [path.name for path in INPUTS if not path.is_file()]
        if missing:
            raise FileNotFoundError(f"缺少原始文件：{missing}")
        input_hashes = {path.name: sha256(path) for path in INPUTS}

        summaries = []
        output_files = []
        for file_number, path in enumerate(INPUTS, 1):
            log(f"[{file_number}/3] 读取 {path.name}")
            raw = read_input(path, file_number)
            log(f"[{file_number}/3] 原始记录 {len(raw):,} 条，开始补点、限幅和长低速处理")
            processed, detail, metrics = preprocess_vehicle_data(raw, config)
            validate_processed_data(processed, config)
            result_path = save_processed(processed, file_number, OUTPUT)
            detail_path = OUTPUT / f"文件{file_number}_处理明细.csv"
            detail.to_csv(detail_path, index=False, encoding="utf-8-sig", date_format="%Y-%m-%d %H:%M:%S")
            metrics["文件"] = f"文件{file_number}"
            summaries.append(metrics)
            output_files.extend([result_path, detail_path])
            log(
                f"[{file_number}/3] 完成：{metrics['原始记录数']:,} → {metrics['处理后记录数']:,} 条；"
                f"补点 {metrics['插值新增数']:,}，修正车速 {metrics['车速修正数']:,}，"
                f"删除长低速记录 {metrics['长低速删除数']:,}"
            )

        summary = pd.DataFrame(summaries)
        summary = summary[["文件"] + [column for column in summary.columns if column != "文件"]]
        summary.to_csv(OUTPUT / "处理汇总.csv", index=False, encoding="utf-8-sig")
        workbook = save_summary(summary, config, OUTPUT)
        figure = save_summary_figure(summary, OUTPUT)
        output_files.extend([OUTPUT / "处理汇总.csv", workbook, figure])
        log(f"汇总图已保存：{figure.relative_to(OUTPUT)}")

        log("[4/4] 重读结果并核验行数、字段、时间和速度约束")
        for file_number, expected in enumerate(summaries, 1):
            path = OUTPUT / f"文件{file_number}_预处理结果.xlsx"
            saved = pd.read_excel(path, sheet_name=f"预处理数据{file_number}")
            if len(saved) != expected["处理后记录数"]:
                raise AssertionError(f"{path.name} 保存行数不一致。")
            validate_processed_data(saved, config)
        saved_summary = pd.read_excel(workbook, sheet_name="处理汇总")
        pd.testing.assert_frame_equal(saved_summary, summary, check_dtype=False)

        for path in INPUTS:
            if sha256(path) != input_hashes[path.name]:
                raise AssertionError(f"原始输入在运行期间发生改变：{path.name}")

        run.update(
            status="completed",
            completed_at=datetime.now().astimezone().isoformat(),
            elapsed_seconds=round(perf_counter() - log.start, 2),
            executable=sys.executable,
            python=platform.python_version(),
            dependencies={name: importlib.metadata.version(name) for name in ["numpy", "pandas", "openpyxl", "matplotlib"]},
            config=config.__dict__,
            input_sha256=input_hashes,
            code_sha256={path.name: sha256(path) for path in [BASE / "main.py", BASE / "preprocessing.py"]},
            output_sha256={str(path.relative_to(OUTPUT)): sha256(path) for path in output_files},
            results=summary.to_dict("records"),
            verification="三份结果重读，时间唯一递增，连续点加减速度、长低速长度和输入哈希全部通过",
        )
        summary_path.write_text(json_text(run), encoding="utf-8")
        log(f"全部完成，结果目录：{OUTPUT}")
    except Exception as exc:
        run.update(status="failed", error=f"{type(exc).__name__}: {exc}")
        summary_path.write_text(json_text(run), encoding="utf-8")
        log(f"运行失败：{exc}")
        raise


if __name__ == "__main__":
    main()
